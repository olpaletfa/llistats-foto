import streamlit as st
import pandas as pd
import json
import os
import requests
import base64
from io import BytesIO
from pdf2image import convert_from_bytes
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from PIL import Image, ImageFile, UnidentifiedImageError, ImageOps
import zipfile

# Librería para generar Excel con imágenes y estilos
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage


# Configuración para permitir imágenes truncadas/incompletas
ImageFile.LOAD_TRUNCATED_IMAGES = True

# --- FUNCIONES DE CONEXIÓN ---

def getToken():
    url = "https://login.salesforce.com/services/oauth2/token"
    payload = {
        "grant_type": "password",
        "client_id": "3MVG9SOw8KERNN08eSQBGpEVIjYdZJXhUjcsFVs3rEm17GWIYlmCs9g0c0.j5YkFKez53A27THurkn2O_T5EP",
        "client_secret": "06D8CB8D5BF09025298677A108257BA9471A6D472EFBC504CCAB6730EE155057",
        "username": "salesforce.api4@escac.es",
        "password": "Fariner@84$",
    }
    headers = {}
    response = requests.request("POST", url, headers=headers, data=payload)
    return response, response.text

def getAlumnes(access_token):
    url = (
        "https://escac4.my.salesforce.com/services/data/v51.0/query/?q="
        "SELECT rio_ed__My_Term__c, Contact_ID_18_Full__c, Grupo__c, rio_ed__Contact_Id__c, "
        "rio_ed__Program_Enrollment__r.Contact_First_Name__c, "
        "rio_ed__Program_Enrollment__r.Contact_Last_Name__c, "
        "rio_ed__Program_Enrollment__r.Contact_Segundo_Apellido__c, " 
        "rio_ed__Program_Enrollment__r.hed__Program_Plan__r.Name, "
        "rio_ed__Program_Enrollment__r.Via_acceso__c, "
        "rio_ed__Program_Enrollment__r.rio_ed__Program_Status__c, "
        "Especialidad__c "
        "FROM rio_ed__PE_Pathway_Status__c "
        "WHERE rio_ed__Term__r.Name = '2026/2027' " 
        "AND ( "
        "rio_ed__Program_Enrollment__r.Tipo_Titulacion__c = 'Grado' OR "
        "rio_ed__Program_Enrollment__r.Tipo_Titulacion__c = 'Curso' OR "
        "rio_ed__Program_Enrollment__r.Tipo_Titulacion__c = 'Máster propio' OR "
        "rio_ed__Program_Enrollment__r.Tipo_Titulacion__c = 'Máster Universitario' OR " 
        "rio_ed__Program_Enrollment__r.Tipo_Titulacion__c = 'Carreras y Estudios Superiores' OR " 
        "rio_ed__Program_Enrollment__r.Tipo_Titulacion__c = 'Postgrado Propio' "
        ") "
        "AND ( "
        "rio_ed__Program_Enrollment__r.rio_ed__Program_Status__c != 'Draft' OR "
        "rio_ed__Program_Enrollment__r.rio_ed__Program_Status__c != 'Cancelled'"
        ") "
        "AND rio_ed__Program_Enrollment__r.Oient_a_Oficial__c = false "
        "AND (rio_ed__Program_Enrollment__r.Via_acceso__c != 'RETITULACIÓ' " 
        "OR rio_ed__Program_Enrollment__r.Via_acceso__c != 'PAS A GRAU')"
    )
    headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
    }
    response = requests.get(url, headers=headers)
    return response, response.text

def getContentDocumentLink(access_token, ContactId):
    url = "https://escac4.my.salesforce.com/services/data/v51.0/query/?q=SELECT ContentDocumentID, LinkedEntityId FROM ContentDocumentLink where LinkedEntityId = '"+ContactId+"'"
    x = {}
    payload = json.dumps(x)
    headers = { "Authorization": "Bearer " + access_token, "Content-Type": "application/json" }
    response = requests.request("GET", url, headers=headers, data=payload)
    resp = json.loads(response.text)
    
    if(resp.get("totalSize", 0) != 0):
        respp = resp["records"]
        for i in respp:
            contentVersion = getConentVersion(access_token, i["ContentDocumentId"])
            if contentVersion is not None:
                return contentVersion 
    return None

def getConentVersion(access_token, DocId):
    url = ("https://escac4.my.salesforce.com/services/data/v51.0/query/?q=SELECT VersionData, FirstPublishLocationId, FileExtension, description, title FROM ContentVersion " 
            "WHERE description = 'Profile picture' AND ContentDocumentId = '"+DocId+"'")
    x = {}
    payload = json.dumps(x)
    headers = { "Authorization": "Bearer " + access_token, "Content-Type": "application/json" }
    response = requests.request("GET", url, headers=headers, data=payload)
    resp = json.loads(response.text)
    if(resp.get("totalSize", 0) != 0):
        return resp["records"][0]["VersionData"]
    return None

def downloadPhoto(access_token, urldata):
    url = "https://escac4.my.salesforce.com" + urldata
    headers = { "Authorization": "Bearer " + access_token }
    response = requests.request("GET", url, headers=headers)
    return response, response.text

def comparar_por_nombre(objeto):
    return objeto.get("nombre_completo", "")

# --- FUNCIÓN GENERACIÓN EXCEL CON FOTOS ---
def makeExcel_Web(lista, dict_fotos):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Listado Alumnos"
    ws.views.sheetView[0].showGridLines = True

    # Estilos de cabecera
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    center_aligned = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_aligned = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers = ["Foto", "Nombre Completo", "Vía de Acceso", "Especialidad", "Term", "Grupo"]
    ws.append(headers)

    # Formatear encabezados
    ws.row_dimensions[1].height = 28
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_aligned

    # Dimensiones de columnas
    ws.column_dimensions['A'].width = 16  # Foto
    ws.column_dimensions['B'].width = 35  # Nombre
    ws.column_dimensions['C'].width = 24  # Vía de Acceso
    ws.column_dimensions['D'].width = 24  # Especialidad
    ws.column_dimensions['E'].width = 14  # Term
    ws.column_dimensions['F'].width = 12  # Grupo

    for idx, alumno in enumerate(lista, start=2):
        ws.row_dimensions[idx].height = 70

        nombre = alumno.get("nombre_completo", "")
        via = alumno.get("Via_acceso__c") or "-"
        especialidad = alumno.get("Especialidad__c") or "-"
        term = str(alumno.get("rio_ed__My_Term__c") or "-")
        grupo = str(alumno.get("Grupo__c") or "-")

        ws.cell(row=idx, column=2, value=nombre)
        ws.cell(row=idx, column=3, value=via)
        ws.cell(row=idx, column=4, value=especialidad)
        ws.cell(row=idx, column=5, value=term)
        ws.cell(row=idx, column=6, value=grupo)

        for col_num in range(1, 7):
            cell = ws.cell(row=idx, column=col_num)
            cell.border = thin_border
            cell.alignment = left_aligned if col_num == 2 else center_aligned

        # Procesar e insertar imagen estandarizada a JPEG
        img_bytes = dict_fotos.get(nombre)
        if img_bytes:
            try:
                with Image.open(BytesIO(img_bytes)) as pil_img:
                    converted_img = pil_img.convert("RGB")
                    
                    # NUEVO: Recortar a cuadrado (60x60) sin deformar. 
                    # centering=(0.5, 0.2) centra horizontalmente, pero recorta más de abajo que de arriba (ideal para caras).
                    converted_img = ImageOps.fit(converted_img, (60, 60), centering=(0.5, 0.2))
                    
                    clean_img_io = BytesIO()
                    converted_img.save(clean_img_io, format="JPEG", quality=85)
                    clean_img_io.seek(0)

                img = OpenpyxlImage(clean_img_io)
                # Ya no forzamos img.width e img.height porque la imagen ya es exactamente de 60x60
                ws.add_image(img, f"A{idx}")
            except Exception:
                ws.cell(row=idx, column=1, value="Sin foto")
        else:
            ws.cell(row=idx, column=1, value="Sin foto")

    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer

# --- FUNCIÓN GENERACIÓN PDF (ROBUSTA) ---
def makePdf_Web(lista, h, nombrePdf, access_token):
    buffer = BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    fuente = "Helvetica"
    tamano_fuente = 20
    c.setFont(fuente, tamano_fuente)
    c.setFillColor("black")
    c.drawString(70, h - 50, nombrePdf)

    # --- Configuración visual ---
    foto_tam = 80          # Tamaño de la foto (ancho y alto)
    salto_y = 95           # Espacio vertical entre cada fila
    items_por_pagina = 7   # Límite por página para evitar que se corte abajo

    # Coordenadas iniciales
    init_ix, init_iy = 100, h - 150
    init_tx, init_ty = 200, h - 115
    init_ax, init_ay = 200, h - 95

    PosIx, PosIy = init_ix, init_iy
    PosTx, PosTy = init_tx, init_ty
    PosAx, PosAy = init_ax, init_ay

    count = 0
    errores_fotos = []

    progress_bar = st.progress(0)
    total = len(lista)

    fotos_para_zip = []
    dict_fotos_bytes = {}

    for idx, j in enumerate(lista):
        progress_bar.progress((idx + 1) / total)
        foto_ok = False

        if "VersionData" in j and j["VersionData"]:
            try:
                _res, _ = downloadPhoto(access_token, j["VersionData"])
                if _res.status_code == 200:
                    image_data = _res.content

                    if image_data[:4] == b'%PDF':
                        errores_fotos.append(f"{j['nombre_completo']} (La foto es un PDF, no imagen)")
                    else:
                        try:
                            # 1. Guardar la imagen original para el ZIP (sin recortar)
                            nombre_limpio = j['nombre_completo'].replace(" ", "_")
                            fotos_para_zip.append({
                                "nombre": f"{nombre_limpio}.jpg",
                                "contenido": image_data
                            })
                            dict_fotos_bytes[j['nombre_completo']] = image_data

                            # 2. Preparar la imagen para el PDF (recortada a cuadrado perfecto)
                            with Image.open(BytesIO(image_data)) as img_pil:
                                img_rgb = img_pil.convert("RGB")
                                # Recortamos a 300x300 (buena resolución para el PDF) enfocando la parte superior
                                img_cropped = ImageOps.fit(img_rgb, (300, 300), centering=(0.5, 0.2))
                                
                                img_stream = BytesIO()
                                img_cropped.save(img_stream, format="JPEG", quality=85)
                                img_stream.seek(0)

                            img_reader = ImageReader(img_stream)

                            # Imagen dibujada en el PDF
                            c.drawImage(img_reader, PosIx, PosIy, width=foto_tam, height=foto_tam, mask='auto')
                            foto_ok = True

                        except UnidentifiedImageError:
                            errores_fotos.append(f"{j['nombre_completo']} (Formato de archivo no reconocido o corrupto)")
                        except Exception as e_img:
                            errores_fotos.append(f"{j['nombre_completo']} (Error interno imagen: {str(e_img)})")
                else:
                    errores_fotos.append(f"{j['nombre_completo']} (Error descarga: {_res.status_code})")

            except Exception as e:
                errores_fotos.append(f"{j['nombre_completo']} (Error general: {e})")
        else:
            errores_fotos.append(f"{j['nombre_completo']} (⚠️ No tiene foto asignada en Salesforce)")

        if not foto_ok:
            c.setStrokeColor("lightgrey")
            c.rect(PosIx, PosIy, foto_tam, foto_tam)
            c.setFillColor("black")

        c.setFont(fuente, 12)
        c.drawString(PosTx, PosTy, j["nombre_completo"])

        if j.get("Via_acceso__c") in ["PAS A GRAU", "TRASLLAT EXPEDIENT", "RETITULACIÓ", "OIENT", "INTERCANVI"]:
            c.setFont(fuente, 8)
            c.drawString(PosAx, PosAy, str(j["Via_acceso__c"]))

        # Desplazamiento vertical entre elementos
        PosAy -= salto_y
        PosTy -= salto_y
        PosIy -= salto_y

        count += 1
        if count == items_por_pagina:
            c.showPage()
            PosIx, PosIy = init_ix, init_iy
            PosTx, PosTy = init_tx, init_ty
            PosAx, PosAy = init_ax, init_ay
            count = 0

    c.showPage()
    c.save()
    progress_bar.empty()
    buffer.seek(0)
    return buffer, errores_fotos, fotos_para_zip, dict_fotos_bytes

# --- APP PRINCIPAL ---
def main():
    st.set_page_config(page_title="Generador Listas ESCAC", layout="wide")
    st.markdown("""<style>#MainMenu {visibility: hidden;} footer {visibility: hidden;}</style>""", unsafe_allow_html=True)
    st.title("🎓 Generador de Listas de Alumnos")

    # 1. Carga de Datos
    if 'data_raw' not in st.session_state:
        with st.spinner('Conectando a Salesforce...'):
            try:
                _res, _text = getToken()
                if _res.status_code != 200:
                    st.error("Error al obtener Token. Revisa credenciales.")
                    st.stop()
                      
                token = json.loads(_text)["access_token"]
                st.session_state.access_token = token
                
                _resAl, _textAl = getAlumnes(token)
                records = json.loads(_textAl)["records"]
                
                clean_data = []
                for obj in records:
                    nuevo = {}
                    nuevo["rio_ed__My_Term__c"] = obj.get("rio_ed__My_Term__c")
                    nuevo["Grupo__c"] = obj.get("Grupo__c")
                    nuevo["Contact_ID_18_Full__c"] = obj.get("Contact_ID_18_Full__c")
                    nuevo["Especialidad__c"] = obj.get("Especialidad__c")
                    pe = obj.get("rio_ed__Program_Enrollment__r", {})
                    if pe:
                        status = pe.get("rio_ed__Program_Status__c", "")
                        if status == "Cancelled":
                            continue
                        nuevo["rio_ed__Program_Status__c"] = pe.get("rio_ed__Program_Status__c", "")
                        nuevo["Contact_First_Name__c"] = pe.get("Contact_First_Name__c", "")
                        nuevo["Contact_Last_Name__c"] = pe.get("Contact_Last_Name__c", "")
                        nuevo["Via_acceso__c"] = pe.get("Via_acceso__c", "")
                        nuevo["nombre_completo"] = f"{nuevo['Contact_First_Name__c']} {nuevo['Contact_Last_Name__c']}"
                        
                        pp = pe.get("hed__Program_Plan__r", {})
                        if pp:
                            raw_name = pp.get("Name", "Desconocido")
                            nombre_lower = raw_name.lower()
                            
                            es_grado = "grado" in nombre_lower and "postgrado" not in nombre_lower
                            es_titulo_propio = "título propio" in nombre_lower
                            
                            if es_grado or es_titulo_propio:
                                nuevo["Name"] = "GRADO"
                                nuevo["Programa_Original"] = raw_name
                            else:
                                nuevo["Name"] = raw_name
                                nuevo["Programa_Original"] = raw_name
                    clean_data.append(nuevo)

                with open('datos_alumnos.json', 'w', encoding='utf-8') as f:
                    json.dump(clean_data, f, ensure_ascii=False, indent=4)
                
                st.session_state.data_raw = clean_data
                st.success("Datos descargados correctamente")
            except Exception as e:
                st.error(f"Error de conexión: {e}")
                st.stop()

    df_alumnos = pd.DataFrame(st.session_state.data_raw)
    access_token = st.session_state.access_token

    # 2. Sidebar
    st.sidebar.header("Filtros")
    programas_disponibles = sorted(df_alumnos["Name"].dropna().unique().tolist())
    curso_seleccionado = st.sidebar.selectbox("Selecciona el Curso / Programa", programas_disponibles)
    
    terms_disponibles = sorted(df_alumnos[df_alumnos["Name"] == curso_seleccionado]["rio_ed__My_Term__c"].dropna().unique().tolist())
    term_seleccionado = st.sidebar.selectbox("Selecciona el Term", terms_disponibles)

    grupos_disponibles = sorted(df_alumnos[
        (df_alumnos["Name"] == curso_seleccionado) & 
        (df_alumnos["rio_ed__My_Term__c"] == term_seleccionado)
    ]["Grupo__c"].dropna().unique().tolist())
    grupo_seleccionado = st.sidebar.selectbox("Selecciona el Grupo", grupos_disponibles)

    # 3. Zona Principal
    st.subheader(f"Listado: {curso_seleccionado}")
    st.info(f"Filtros: Term {term_seleccionado} | Grupo {grupo_seleccionado}")

    df_filtrado = df_alumnos[
        (df_alumnos["Name"] == curso_seleccionado) &
        (df_alumnos["rio_ed__My_Term__c"] == term_seleccionado) &
        (df_alumnos["Grupo__c"] == grupo_seleccionado)
    ].copy()

    st.write(f"Alumnos encontrados: **{len(df_filtrado)}**")
    
    st.dataframe(df_filtrado[["nombre_completo", "Via_acceso__c", "Especialidad__c"]], width="stretch")

    # 4. Botón Generar
    if len(df_filtrado) > 0:
        if st.button("📸 Procesar Fotos y Generar Documentos"):
            lista_para_procesar = df_filtrado.to_dict('records')
            
            with st.spinner("Buscando fotos en Salesforce..."):
                lista_final = []
                progress_text = st.empty()
                
                for idx, alumno in enumerate(lista_para_procesar):
                    progress_text.text(f"Obteniendo foto {idx+1}/{len(lista_para_procesar)}: {alumno['nombre_completo']}")
                    contentVersionUrl = getContentDocumentLink(access_token, alumno["Contact_ID_18_Full__c"])
                    alumno["VersionData"] = contentVersionUrl
                    lista_final.append(alumno)
                
                progress_text.empty()
                lista_final = sorted(lista_final, key=comparar_por_nombre)
                
                # Generar PDF y capturar fotos en memoria
                h_page = A4[1]
                nombre_base_archivo = f"{curso_seleccionado}_T{term_seleccionado}_G{grupo_seleccionado}".replace(" ", "_")
                
                pdf_bytes, lista_errores, fotos_para_zip, dict_fotos = makePdf_Web(
                    lista_final, h_page, nombre_base_archivo, access_token
                )
                
                # Generar Excel con fotos incrustadas
                excel_bytes = makeExcel_Web(lista_final, dict_fotos)
                
                json_string = json.dumps(lista_final, indent=4, ensure_ascii=False)
                
                st.success("¡Proceso finalizado con éxito!")
                
                # Avisos de fotos
                if len(lista_errores) > 0:
                    with st.expander(f"⚠️ Avisos de Fotos ({len(lista_errores)} alumnos)", expanded=True):
                        st.warning("Los siguientes alumnos no tienen foto o dio error al cargarla:")
                        for err in lista_errores:
                            st.write(f"- {err}")
                            
                # Opciones de Descarga
                st.markdown("### 📥 Descargas disponibles")
                col1, col2, col3, col4 = st.columns(4)

                with col1:
                    st.download_button(
                        label="⬇️ Descargar PDF",
                        data=pdf_bytes,
                        file_name=f"{nombre_base_archivo}.pdf",
                        mime="application/pdf",
                        use_container_width=True
                    )

                with col2:
                    st.download_button(
                        label="📊 Descargar Excel",
                        data=excel_bytes,
                        file_name=f"{nombre_base_archivo}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True
                    )

                with col3:
                    if fotos_para_zip:
                        zip_buffer = BytesIO()
                        with zipfile.ZipFile(zip_buffer, "w") as zf:
                            for f in fotos_para_zip:
                                zf.writestr(f["nombre"], f["contenido"])
                        zip_buffer.seek(0)
                        st.download_button(
                            label="🗂️ Descargar Fotos (ZIP)",
                            data=zip_buffer,
                            file_name=f"{nombre_base_archivo}_fotos.zip",
                            mime="application/zip",
                            use_container_width=True
                        )
                    else:
                        st.button("🗂️ Sin fotos para ZIP", disabled=True, use_container_width=True)

                with col4:
                    st.download_button(
                        label="💾 Descargar JSON",
                        data=json_string,
                        file_name=f"{nombre_base_archivo}.json",
                        mime="application/json",
                        use_container_width=True
                    )

if __name__ == "__main__":
    main()